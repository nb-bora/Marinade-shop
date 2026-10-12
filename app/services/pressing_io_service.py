"""Import et modèle Excel/CSV des articles et combinaisons d'un pressing.

Réutilise la mécanique de ``catalog_io_service`` (lecture CSV ``;``, validation de cellules,
résultat ``{imported, created, updated, errors}``, une ligne = un point de sauvegarde, donc
import partiel) avec les mêmes limites : 5 Mo, 5000 lignes.

Colonnes
--------
* ``articles`` : ``nom;prix_unitaire;categorie``
* ``combinaisons`` : ``nom;prix;article_1;quantite_1;...;article_5;quantite_5``
  (``article_n`` = nom d'un article ACTIF du restaurant ; ``quantite_n`` vaut 1 si vide)

Idempotent par nom (sans la casse) : une ligne dont le nom existe met l'élément à jour (le
prix, la catégorie ; pour une combinaison, ses lignes sont remplacées). L'état actif/inactif
n'est jamais modifié par un import. Prix et quantités sont des entiers.
"""

from __future__ import annotations

import io
import uuid
from decimal import Decimal
from typing import Callable, Dict, List, Tuple

from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.models.pressing import PressingArticle
from app.services.catalog_io_service import (
    MAX_FILE_BYTES,
    MAX_ROWS,
    FileFormatError,
    ImportResult,
    Row,
    RowError,
    _cell_text,
    _check_header,
    _decimal,
    _text,
    read_csv_rows,
)
from app.services.pressing_calcul import PressingRegleError
from app.services.pressing_service import PRIX_MAX, PressingService
from app.utils.exceptions import ConflictError
from app.utils.logging import get_logger

logger = get_logger(__name__)

MAX_ARTICLES_PAR_COMBINAISON = 5
ARTICLES_HEADER = ["nom", "prix_unitaire", "categorie"]
COMBINAISONS_HEADER = ["nom", "prix"] + [
    colonne
    for i in range(1, MAX_ARTICLES_PAR_COMBINAISON + 1)
    for colonne in (f"article_{i}", f"quantite_{i}")
]
HEADERS = {"articles": ARTICLES_HEADER, "combinaisons": COMBINAISONS_HEADER}
SHEET_NAMES = {"articles": "Articles", "combinaisons": "Combinaisons"}
IMPORT_TYPES = tuple(HEADERS)
ACTION_IMPORT = "PRESSING_IMPORT"


def template_xlsx(import_type: str) -> bytes:
    """Classeur vierge : une feuille, en-têtes seulement."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    header = HEADERS.get(import_type)
    if header is None:
        raise FileFormatError("type de modèle inconnu : " + ", ".join(IMPORT_TYPES))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAMES[import_type]
    sheet.append(header)
    for index, cell in enumerate(sheet[1], start=1):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="EA6A2B")
        sheet.column_dimensions[cell.column_letter].width = max(14, len(header[index - 1]) + 6)
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_xlsx_rows(content: bytes, import_type: str) -> Tuple[List[str], List[Row]]:
    """Feuille « Articles »/« Combinaisons », à défaut la première feuille du classeur."""
    from openpyxl import load_workbook

    if len(content) > MAX_FILE_BYTES:
        raise FileFormatError("Fichier trop volumineux (5 Mo maximum)")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise FileFormatError("Fichier XLSX illisible") from exc
    try:
        wanted = SHEET_NAMES[import_type].lower()
        sheets = workbook.worksheets
        sheet = next((s for s in sheets if s.title.strip().lower() == wanted), None)
        sheet = sheet or (sheets[0] if sheets else None)
        if sheet is None:
            raise FileFormatError("Classeur vide")
        iterator = sheet.iter_rows(values_only=True)
        first = next(iterator, None)
        if first is None:
            raise FileFormatError("Fichier vide")
        header = [("" if c is None else str(c)).strip().lower() for c in first]
        rows: List[Row] = []
        for index, values in enumerate(iterator, start=2):
            cells = ["" if v is None else _cell_text(v) for v in values]
            if not any(cells):
                continue
            if len(rows) >= MAX_ROWS:
                raise FileFormatError(f"Trop de lignes ({MAX_ROWS} maximum par import)")
            rows.append(
                (index, {h: (cells[i] if i < len(cells) else "") for i, h in enumerate(header) if h})
            )
        return header, rows
    finally:
        workbook.close()


def read_rows(content: bytes, file_format: str, import_type: str):
    if file_format == "csv":
        return read_csv_rows(content)
    return read_xlsx_rows(content, import_type)


def _entier(row: Dict[str, str], name: str, *, required: bool = False, minimum: int = 0):
    value = _decimal(
        row, name, required=required, minimum=Decimal(minimum), maximum=Decimal(PRIX_MAX)
    )
    if value is None:
        return None
    if value != value.to_integral_value():
        raise RowError(name, "Doit être un nombre entier (FCFA)")
    return int(value)


class PressingImporter:
    def __init__(self, db: Session, restaurant_id: uuid.UUID, actor_id=None):
        self.db = db
        self.service = PressingService(db, restaurant_id, actor_id)
        self.result = ImportResult()
        self._articles: Dict[str, PressingArticle] | None = None

    @property
    def articles(self) -> Dict[str, PressingArticle]:
        if self._articles is None:
            self._articles = {a.nom.lower(): a for a in self.service.list_articles()}
        return self._articles

    def run(self, import_type: str, header: List[str], rows: List[Row]) -> ImportResult:
        if import_type == "articles":
            _check_header(header, ["nom", "prix_unitaire"], ARTICLES_HEADER)
            handler: Callable = self._article_row
        else:
            _check_header(header, ["nom", "prix", "article_1"], COMBINAISONS_HEADER)
            handler = self._combinaison_row
        for row_number, data in rows:
            self._run_row(row_number, data, handler)
        return self.result

    def _run_row(self, row_number: int, data: Dict[str, str], handler: Callable) -> None:
        try:
            with self.db.begin_nested():
                created = handler(data)
        except RowError as exc:
            self.result.add_error(row_number, exc.field, exc.message)
            self._articles = None
        except (ConflictError, PressingRegleError) as exc:
            self.result.add_error(row_number, "nom", exc.message)
            self._articles = None
        except SQLAlchemyError:
            logger.exception("Import pressing : ligne %s refusée par la base", row_number)
            self.result.add_error(row_number, "", "Ligne refusée par la base de données")
            self._articles = None
        else:
            self.result.imported += 1
            if created:
                self.result.created += 1
            else:
                self.result.updated += 1

    def _article_row(self, row: Dict[str, str]) -> bool:
        nom = _text(row, "nom", required=True, max_len=120)
        prix = _entier(row, "prix_unitaire", required=True)
        categorie = _text(row, "categorie", max_len=120)
        existing = self.articles.get(nom.lower())
        if existing is None:
            self.articles[nom.lower()] = self.service.create_article(nom, prix, categorie)
            return True
        self.service.update_article(
            existing.id, existing.nom, prix,
            categorie if categorie is not None else existing.categorie,
            existing.actif,
        )  # fmt: skip
        return False

    def _combinaison_row(self, row: Dict[str, str]) -> bool:
        nom = _text(row, "nom", required=True, max_len=120)
        prix = _entier(row, "prix", required=True)
        lignes: List[tuple] = []
        for i in range(1, MAX_ARTICLES_PAR_COMBINAISON + 1):
            nom_article = _text(row, f"article_{i}", max_len=120)
            quantite = _entier(row, f"quantite_{i}", minimum=1)
            if nom_article is None:
                if quantite is not None:
                    raise RowError(f"article_{i}", "Article manquant pour cette quantité")
                continue
            article = self.articles.get(nom_article.lower())
            if article is None:
                raise RowError(f"article_{i}", f"Article inconnu : « {nom_article} »")
            if not article.actif:
                raise RowError(f"article_{i}", f"Article désactivé : « {nom_article} »")
            lignes.append((article.id, quantite or 1))
        if not lignes:
            raise RowError("article_1", "Au moins un article est obligatoire")
        if len({article_id for article_id, _ in lignes}) != len(lignes):
            raise RowError("article_1", "Un article ne peut apparaître qu'une fois")
        existing = self.service._combinaison_by_name(nom)
        if existing is None:
            self.service.create_combinaison(nom, prix, lignes)
            return True
        self.service.update_combinaison(existing.id, existing.nom, prix, lignes, existing.actif)
        return False


def import_rows(
    db: Session, restaurant_id: uuid.UUID, actor_id, content: bytes, file_format: str,
    import_type: str,
) -> ImportResult:
    if import_type not in HEADERS:
        raise FileFormatError("type d'import inconnu : " + ", ".join(IMPORT_TYPES))
    header, rows = read_rows(content, file_format, import_type)
    return PressingImporter(db, restaurant_id, actor_id).run(import_type, header, rows)
