"""Import et export du catalogue d'un restaurant (plats, boissons, composants, recettes).

Formats
-------
CSV : séparateur ``;``, UTF-8 (BOM accepté ; cp1252 toléré en lecture), première ligne
d'en-tête. Un fichier ne porte qu'UN type d'import :

* ``catalogue`` (plats ET boissons, la colonne ``type`` aiguille chaque ligne) et
  ``plats`` / ``boissons`` : ``type;nom;prix;description;code;disponible;stock_min;stock_max;composition``
* ``composants`` : ``nom;prix;stock;composition;valeur_calorique`` (``type`` facultatif)
* ``recettes`` : ``plat;composant;quantite;unite``

XLSX : une feuille par type (« Plats », « Boissons », « Composants », « Recettes »), mêmes
colonnes. Les feuilles absentes sont ignorées.

Règles
------
* Import idempotent : une ligne dont le nom existe déjà MET À JOUR l'élément (clé :
  restaurant + nom, sans tenir compte de la casse) ; ré-importer un fichier ne duplique rien.
* Import partiel : chaque ligne s'exécute dans son propre point de sauvegarde ; une ligne
  invalide est signalée ``{row, field, message}`` sans emporter les autres.
* ``composition`` d'un plat : ``Riz|Poulet:2`` (noms de composants du restaurant, quantité
  facultative après ``:``). Pour un composant c'est un texte libre (ingrédients).
* ``code``, ``stock_min`` et ``stock_max`` des plats/boissons sont conservés dans
  ``attributs_jsonb`` ; ``valeur_calorique`` et ``composition`` des composants aussi.
"""

from __future__ import annotations

import csv
import io
import uuid
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple

from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.models.restaurant import (
    Boisson,
    Composant,
    Plat,
    PlatComposant,
    StockComposant,
    StockMouvement,
)
from app.utils.logging import get_logger

logger = get_logger(__name__)

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000

CATALOG_HEADER = [
    "type", "nom", "prix", "description", "code", "disponible",
    "stock_min", "stock_max", "composition",
]  # fmt: skip
COMPOSANT_HEADER = ["nom", "prix", "stock", "composition", "valeur_calorique"]
RECETTE_HEADER = ["plat", "composant", "quantite", "unite"]

SHEETS = {
    "Plats": CATALOG_HEADER,
    "Boissons": CATALOG_HEADER,
    "Composants": COMPOSANT_HEADER,
    "Recettes": RECETTE_HEADER,
}
# Ordre d'import : les recettes et compositions référencent des composants et des plats.
SHEET_ORDER = ["Composants", "Plats", "Boissons", "Recettes"]
IMPORT_TYPES = {"catalogue", "plats", "boissons", "composants", "recettes"}

_MAX_PRICE = Decimal("99999999.99")
_TRUE = {"1", "true", "vrai", "oui", "o", "yes", "y", "x"}
_FALSE = {"0", "false", "faux", "non", "n", "no"}


class FileFormatError(ValueError):
    """Le fichier lui-même est inutilisable (encodage, en-tête, taille) : HTTP 400."""


class RowError(Exception):
    def __init__(self, field_name: str, message: str):
        super().__init__(message)
        self.field = field_name
        self.message = message


@dataclass
class ImportResult:
    imported: int = 0
    created: int = 0
    updated: int = 0
    errors: List[Dict[str, Any]] = field(default_factory=list)

    def add_error(self, row: int, field_name: str, message: str, sheet: str | None = None):
        error: Dict[str, Any] = {"row": row, "field": field_name, "message": message}
        if sheet:
            error["sheet"] = sheet
        self.errors.append(error)

    def as_dict(self) -> dict:
        return {
            "imported": self.imported,
            "created": self.created,
            "updated": self.updated,
            "errors": self.errors,
        }


# --------------------------------------------------------------------------- #
# Lecture des fichiers -> lignes {colonne: texte}
# --------------------------------------------------------------------------- #
Row = Tuple[int, Dict[str, str]]


def _check_header(found: List[str], required: Iterable[str], allowed: Iterable[str]) -> None:
    missing = [c for c in required if c not in found]
    if missing:
        raise FileFormatError(
            "Colonnes manquantes : " + ", ".join(missing) + ". En-tête attendu : "
            + ";".join(allowed)
        )


def read_csv_rows(content: bytes) -> Tuple[List[str], List[Row]]:
    if len(content) > MAX_FILE_BYTES:
        raise FileFormatError("Fichier trop volumineux (5 Mo maximum)")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise FileFormatError("Encodage illisible : enregistrez en UTF-8") from exc
    reader = csv.reader(io.StringIO(text), delimiter=";")
    try:
        header_raw = next(reader)
    except StopIteration:
        raise FileFormatError("Fichier vide") from None
    except csv.Error as exc:
        raise FileFormatError(f"CSV illisible : {exc}") from exc
    header = [h.strip().lower() for h in header_raw]
    if len(header) == 1 and ("," in header[0] or "\t" in header[0]):
        raise FileFormatError("Séparateur invalide : utilisez le point-virgule (;)")
    rows: List[Row] = []
    try:
        for index, values in enumerate(reader, start=2):
            if not any(v.strip() for v in values):
                continue
            if len(rows) >= MAX_ROWS:
                raise FileFormatError(f"Trop de lignes ({MAX_ROWS} maximum par import)")
            rows.append(
                (index, {h: (values[i].strip() if i < len(values) else "") for i, h in enumerate(header) if h})
            )
    except csv.Error as exc:
        raise FileFormatError(f"CSV illisible : {exc}") from exc
    return header, rows


def read_xlsx_sheets(content: bytes) -> Dict[str, Tuple[List[str], List[Row]]]:
    from openpyxl import load_workbook

    if len(content) > MAX_FILE_BYTES:
        raise FileFormatError("Fichier trop volumineux (5 Mo maximum)")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # fichier corrompu, mauvais format, zip invalide…
        raise FileFormatError("Fichier XLSX illisible") from exc
    sheets: Dict[str, Tuple[List[str], List[Row]]] = {}
    total = 0
    try:
        wanted = {name.lower(): name for name in SHEETS}
        for ws in workbook.worksheets:
            name = wanted.get(ws.title.strip().lower())
            if name is None:
                continue
            iterator = ws.iter_rows(values_only=True)
            first = next(iterator, None)
            if first is None:
                continue
            header = [("" if c is None else str(c)).strip().lower() for c in first]
            rows: List[Row] = []
            for index, values in enumerate(iterator, start=2):
                cells = ["" if v is None else _cell_text(v) for v in values]
                if not any(c for c in cells):
                    continue
                total += 1
                if total > MAX_ROWS:
                    raise FileFormatError(f"Trop de lignes ({MAX_ROWS} maximum par import)")
                rows.append(
                    (index, {h: (cells[i] if i < len(cells) else "") for i, h in enumerate(header) if h})
                )
            sheets[name] = (header, rows)
    finally:
        workbook.close()
    return sheets


def _cell_text(value: Any) -> str:
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# --------------------------------------------------------------------------- #
# Validation des cellules
# --------------------------------------------------------------------------- #
def _text(row: Dict[str, str], name: str, *, required: bool = False, max_len: int = 255):
    value = row.get(name, "").strip()
    if not value:
        if required:
            raise RowError(name, "Valeur obligatoire")
        return None
    if len(value) > max_len:
        raise RowError(name, f"{max_len} caractères maximum")
    return value


def _decimal(
    row: Dict[str, str], name: str, *, required: bool = False, minimum: Decimal = Decimal(0),
    strictly_positive: bool = False, maximum: Decimal = _MAX_PRICE,
) -> Optional[Decimal]:  # fmt: skip
    raw = row.get(name, "").strip()
    if not raw:
        if required:
            raise RowError(name, "Valeur obligatoire")
        return None
    cleaned = raw.replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        raise RowError(name, f"Nombre invalide : « {raw} »") from None
    if not value.is_finite():
        raise RowError(name, f"Nombre invalide : « {raw} »")
    if strictly_positive and value <= 0:
        raise RowError(name, "Doit être strictement positif")
    if value < minimum:
        raise RowError(name, f"Doit être supérieur ou égal à {minimum}")
    if value > maximum:
        raise RowError(name, "Valeur trop grande")
    return value


def _bool(row: Dict[str, str], name: str) -> Optional[bool]:
    raw = row.get(name, "").strip().lower()
    if not raw:
        return None
    if raw in _TRUE:
        return True
    if raw in _FALSE:
        return False
    raise RowError(name, "Attendu : oui/non (ou true/false, 1/0)")


def _parse_composition(raw: str) -> List[Tuple[str, Decimal]]:
    """``Riz|Poulet:2`` -> [("Riz", 1), ("Poulet", 2)]."""
    lines: List[Tuple[str, Decimal]] = []
    seen = set()
    for chunk in (c.strip() for c in raw.split("|")):
        if not chunk:
            continue
        name, _, qty = chunk.partition(":")
        name = name.strip()
        try:
            quantity = Decimal(qty.strip().replace(",", ".")) if qty.strip() else Decimal(1)
        except InvalidOperation:
            raise RowError("composition", f"Quantité invalide pour « {name} »") from None
        if not name or quantity <= 0 or quantity > Decimal("999999999"):
            raise RowError("composition", f"Entrée invalide : « {chunk} »")
        if name.lower() in seen:
            raise RowError("composition", f"Composant en double : « {name} »")
        seen.add(name.lower())
        lines.append((name, quantity))
    return lines


# --------------------------------------------------------------------------- #
# Import
# --------------------------------------------------------------------------- #
class CatalogImporter:
    def __init__(self, db: Session, restaurant_id: uuid.UUID):
        self.db = db
        self.restaurant_id = restaurant_id
        self.result = ImportResult()
        self._plats: Dict[str, Plat] | None = None
        self._boissons: Dict[str, Boisson] | None = None
        self._composants: Dict[str, Composant] | None = None

    # -- caches par nom (une requête par table, pas par ligne) ------------------
    def _by_name(self, model) -> Dict[str, Any]:
        rows = (
            self.db.query(model)
            .filter(model.restaurant_id == self.restaurant_id, model.deleted_at.is_(None))
            .order_by(model.created_at, model.id)
            .all()
        )
        return {r.nom.lower(): r for r in rows}

    @property
    def plats(self) -> Dict[str, Plat]:
        if self._plats is None:
            self._plats = self._by_name(Plat)
        return self._plats

    @property
    def boissons(self) -> Dict[str, Boisson]:
        if self._boissons is None:
            self._boissons = self._by_name(Boisson)
        return self._boissons

    @property
    def composants(self) -> Dict[str, Composant]:
        if self._composants is None:
            self._composants = self._by_name(Composant)
        return self._composants

    # -- exécution ----------------------------------------------------------------
    def run(self, import_type: str, header: List[str], rows: List[Row], sheet: str | None = None):
        if import_type in ("catalogue", "plats", "boissons"):
            required = ["nom", "prix"] + (["type"] if import_type == "catalogue" else [])
            _check_header(header, required, CATALOG_HEADER)
            handler = lambda r: self._catalog_row(import_type, r)  # noqa: E731
        elif import_type == "composants":
            _check_header(header, ["nom"], COMPOSANT_HEADER)
            handler = self._composant_row
        else:
            _check_header(header, ["plat", "composant"], RECETTE_HEADER)
            handler = self._recette_row
        for row_number, data in rows:
            self._run_row(row_number, data, handler, sheet)

    def _run_row(self, row_number: int, data: Dict[str, str], handler: Callable, sheet):
        try:
            with self.db.begin_nested():
                created = handler(data)
        except RowError as exc:
            self.result.add_error(row_number, exc.field, exc.message, sheet)
            self._reset_caches()
        except SQLAlchemyError:
            logger.exception("Catalog import: row %s rejected by the database", row_number)
            self.result.add_error(row_number, "", "Ligne refusée par la base de données", sheet)
            self._reset_caches()
        else:
            self.result.imported += 1
            if created:
                self.result.created += 1
            else:
                self.result.updated += 1

    def _reset_caches(self) -> None:
        # Un point de sauvegarde annulé peut laisser des objets du cache détachés
        # de la base : on relit plutôt que d'en réutiliser un état incertain.
        self._plats = self._boissons = self._composants = None

    # -- lignes -------------------------------------------------------------------
    def _catalog_row(self, import_type: str, row: Dict[str, str]) -> bool:
        kind = row.get("type", "").strip().lower()
        expected = {"plats": "plat", "boissons": "boisson"}.get(import_type)
        if kind in ("plats", "boissons"):
            kind = kind[:-1]
        if import_type == "catalogue":
            if kind not in ("plat", "boisson"):
                raise RowError("type", "Attendu : plat ou boisson")
        else:
            if kind and kind != expected:
                raise RowError("type", f"Cette feuille/import ne contient que des {expected}s")
            kind = expected

        nom = _text(row, "nom", required=True)
        prix = _decimal(row, "prix", required=True, strictly_positive=True)
        description = _text(row, "description", max_len=2000)
        code = _text(row, "code", max_len=50)
        disponible = _bool(row, "disponible")
        stock_min = _decimal(row, "stock_min", maximum=Decimal("999999999"))
        stock_max = _decimal(row, "stock_max", maximum=Decimal("999999999"))
        if stock_min is not None and stock_max is not None and stock_min > stock_max:
            raise RowError("stock_min", "stock_min ne peut pas dépasser stock_max")
        composition = row.get("composition", "").strip()
        if kind == "boisson" and composition:
            raise RowError("composition", "Une boisson n'a pas de composition")
        lines = _parse_composition(composition) if composition else []
        resolved: List[Tuple[Composant, Decimal]] = []
        for name, quantity in lines:
            composant = self.composants.get(name.lower())
            if composant is None:
                raise RowError("composition", f"Composant inconnu : « {name} »")
            resolved.append((composant, quantity))

        extras = {
            k: str(v) if isinstance(v, Decimal) else v
            for k, v in (("code", code), ("stock_min", stock_min), ("stock_max", stock_max))
            if v is not None
        }
        cache = self.plats if kind == "plat" else self.boissons
        existing = cache.get(nom.lower())
        created = existing is None
        if created:
            model = Plat if kind == "plat" else Boisson
            existing = model(
                restaurant_id=self.restaurant_id, nom=nom, prix=prix, devise="XAF",
                disponible=True if disponible is None else disponible,
                description=description, attributs_jsonb=extras or None,
            )  # fmt: skip
            self.db.add(existing)
            self.db.flush()
            cache[nom.lower()] = existing
        else:
            existing.prix = prix
            if description is not None:
                existing.description = description
            if disponible is not None:
                existing.disponible = disponible
            if extras:
                existing.attributs_jsonb = {**(existing.attributs_jsonb or {}), **extras}
        if kind == "plat" and composition:
            self.db.query(PlatComposant).filter(PlatComposant.plat_id == existing.id).delete(
                synchronize_session=False
            )
            for order, (composant, quantity) in enumerate(resolved):
                self.db.add(
                    PlatComposant(
                        plat_id=existing.id, composant_id=composant.id,
                        quantite=quantity, unite=composant.stock_unite or "portion", ordre=order,
                    )
                )  # fmt: skip
        self.db.flush()
        return created

    def _composant_row(self, row: Dict[str, str]) -> bool:
        nom = _text(row, "nom", required=True)
        prix = _decimal(row, "prix")
        stock = _decimal(row, "stock", maximum=Decimal("999999999"))
        composition = _text(row, "composition", max_len=2000)
        calories = _decimal(row, "valeur_calorique", maximum=Decimal("99999"))
        kind = _text(row, "type", max_len=30)
        extras = {}
        if composition is not None:
            extras["composition"] = composition
        if calories is not None:
            extras["valeur_calorique"] = str(calories)

        existing = self.composants.get(nom.lower())
        created = existing is None
        if created:
            existing = Composant(
                restaurant_id=self.restaurant_id, nom=nom, type=kind or "ingredient",
                prix_supplement=prix or Decimal(0), devise="XAF",
                attributs_jsonb=extras or None,
            )  # fmt: skip
            self.db.add(existing)
            self.db.flush()
            self.db.add(StockComposant(composant_id=existing.id, quantite=Decimal(0)))
            self.db.flush()  # la session n'autoflush pas : _set_stock doit le retrouver
            self.composants[nom.lower()] = existing
        else:
            if prix is not None:
                existing.prix_supplement = prix
            if kind:
                existing.type = kind
            if extras:
                existing.attributs_jsonb = {**(existing.attributs_jsonb or {}), **extras}
        if stock is not None:
            self._set_stock(existing, stock)
        self.db.flush()
        return created

    def _set_stock(self, composant: Composant, quantity: Decimal) -> None:
        stock = (
            self.db.query(StockComposant)
            .filter(StockComposant.composant_id == composant.id)
            .with_for_update()
            .first()
        )
        if stock is None:
            stock = StockComposant(composant_id=composant.id, quantite=Decimal(0))
            self.db.add(stock)
            self.db.flush()
        if quantity < stock.reservee:
            raise RowError("stock", "Inférieur à la quantité déjà réservée par des commandes")
        if quantity == stock.quantite:
            return
        movement = "entree" if quantity > stock.quantite else "ajustement"
        amount = quantity - stock.quantite if movement == "entree" else quantity
        stock.quantite = quantity
        if amount > 0:
            self.db.add(
                StockMouvement(
                    composant_id=composant.id, type=movement, quantite=amount,
                    reference_type="import", notes="Import de catalogue",
                )
            )  # fmt: skip

    def _recette_row(self, row: Dict[str, str]) -> bool:
        plat_name = _text(row, "plat", required=True)
        composant_name = _text(row, "composant", required=True)
        quantite = _decimal(
            row, "quantite", strictly_positive=True, maximum=Decimal("999999999")
        ) or Decimal(1)
        unite = _text(row, "unite", max_len=20)
        plat = self.plats.get(plat_name.lower())
        if plat is None:
            raise RowError("plat", f"Plat inconnu : « {plat_name} »")
        composant = self.composants.get(composant_name.lower())
        if composant is None:
            raise RowError("composant", f"Composant inconnu : « {composant_name} »")
        line = (
            self.db.query(PlatComposant)
            .filter(
                PlatComposant.plat_id == plat.id,
                PlatComposant.composant_id == composant.id,
            )
            .first()
        )
        if line is not None:
            line.quantite = quantite
            if unite:
                line.unite = unite
            self.db.flush()
            return False
        order = (
            self.db.query(PlatComposant).filter(PlatComposant.plat_id == plat.id).count()
        )
        self.db.add(
            PlatComposant(
                plat_id=plat.id, composant_id=composant.id, quantite=quantite,
                unite=unite or composant.stock_unite or "portion", ordre=order,
            )
        )  # fmt: skip
        self.db.flush()
        return True


def import_csv(db: Session, restaurant_id: uuid.UUID, content: bytes, import_type: str) -> ImportResult:
    if import_type not in IMPORT_TYPES:
        raise FileFormatError("type d'import inconnu : " + ", ".join(sorted(IMPORT_TYPES)))
    header, rows = read_csv_rows(content)
    importer = CatalogImporter(db, restaurant_id)
    importer.run(import_type, header, rows)
    return importer.result


def import_xlsx(db: Session, restaurant_id: uuid.UUID, content: bytes) -> ImportResult:
    sheets = read_xlsx_sheets(content)
    if not sheets:
        raise FileFormatError(
            "Aucune feuille reconnue : attendu « Plats », « Boissons », « Composants » ou « Recettes »"
        )
    importer = CatalogImporter(db, restaurant_id)
    kinds = {"Plats": "plats", "Boissons": "boissons", "Composants": "composants", "Recettes": "recettes"}
    for name in SHEET_ORDER:
        if name in sheets:
            header, rows = sheets[name]
            importer.run(kinds[name], header, rows, sheet=name)
    return importer.result


# --------------------------------------------------------------------------- #
# Export
# --------------------------------------------------------------------------- #
def _safe_cell(value: Any) -> Any:
    """Neutralise l'injection de formules : Excel exécute une cellule texte qui commence par = + - @."""
    if isinstance(value, str) and value and value[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _number(value: Any) -> Any:
    return None if value is None else float(value)


def export_xlsx(db: Session, restaurant_id: uuid.UUID) -> bytes:
    from openpyxl import Workbook

    composants = (
        db.query(Composant)
        .filter(Composant.restaurant_id == restaurant_id, Composant.deleted_at.is_(None))
        .order_by(Composant.nom, Composant.id)
        .all()
    )
    stocks = {
        s.composant_id: s
        for s in db.query(StockComposant)
        .filter(StockComposant.composant_id.in_([c.id for c in composants] or [None]))
        .all()
    }
    plats = (
        db.query(Plat)
        .filter(Plat.restaurant_id == restaurant_id, Plat.deleted_at.is_(None))
        .order_by(Plat.nom, Plat.id)
        .all()
    )
    boissons = (
        db.query(Boisson)
        .filter(Boisson.restaurant_id == restaurant_id, Boisson.deleted_at.is_(None))
        .order_by(Boisson.nom, Boisson.id)
        .all()
    )
    lines = (
        db.query(PlatComposant)
        .filter(PlatComposant.plat_id.in_([p.id for p in plats] or [None]))
        .order_by(PlatComposant.plat_id, PlatComposant.ordre)
        .all()
    )
    by_plat: Dict[uuid.UUID, List[PlatComposant]] = {}
    for line in lines:
        by_plat.setdefault(line.plat_id, []).append(line)
    comp_name = {c.id: c.nom for c in composants}
    plat_name = {p.id: p.nom for p in plats}

    def catalog_row(kind: str, item, composition: str = "") -> list:
        extra = item.attributs_jsonb or {}
        return [
            kind, _safe_cell(item.nom), _number(item.prix), _safe_cell(item.description),
            _safe_cell(extra.get("code")), "oui" if item.disponible else "non",
            _number(Decimal(extra["stock_min"])) if extra.get("stock_min") is not None else None,
            _number(Decimal(extra["stock_max"])) if extra.get("stock_max") is not None else None,
            _safe_cell(composition or None),
        ]  # fmt: skip

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Plats")
    sheet.append(CATALOG_HEADER)
    for p in plats:
        composition = "|".join(
            f"{comp_name.get(l.composant_id, '?')}:{l.quantite.normalize():f}"
            for l in by_plat.get(p.id, [])
        )
        sheet.append(catalog_row("plat", p, composition))
    sheet = workbook.create_sheet("Boissons")
    sheet.append(CATALOG_HEADER)
    for b in boissons:
        sheet.append(catalog_row("boisson", b))
    sheet = workbook.create_sheet("Composants")
    sheet.append(COMPOSANT_HEADER)
    for c in composants:
        extra = c.attributs_jsonb or {}
        stock = stocks.get(c.id)
        sheet.append(
            [
                _safe_cell(c.nom), _number(c.prix_supplement),
                _number(stock.quantite) if stock else 0,
                _safe_cell(extra.get("composition")),
                _number(Decimal(extra["valeur_calorique"])) if extra.get("valeur_calorique") else None,
            ]
        )  # fmt: skip
    sheet = workbook.create_sheet("Recettes")
    sheet.append(RECETTE_HEADER)
    for line in lines:
        sheet.append(
            [
                _safe_cell(plat_name.get(line.plat_id)),
                _safe_cell(comp_name.get(line.composant_id)),
                _number(line.quantite), _safe_cell(line.unite),
            ]
        )  # fmt: skip
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
