"""Tarification d'un pressing : calcul (unitaire) et API (intégration PostgreSQL).

Les tests unitaires n'ont besoin d'aucune base. Les tests d'intégration exigent
TEST_DATABASE_URL (voir conftest.py) :

    pytest tests/test_pressing.py -v
"""

import io
import uuid
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from app.services.pressing_calcul import (
    ArticleTarifie,
    CombinaisonTarifee,
    DemandeCalcul,
    PressingRegleError,
    Tarif,
    calculer,
    montant_au_kilo,
    poids_facturable,
)
from tests.support import headers_for, make_member, register

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV = "text/csv"


# --------------------------------------------------------------------------- #
# Unitaires : service de calcul pur
# --------------------------------------------------------------------------- #
TARIF = Tarif(prix_kg=500, prix_sechage_kg=None, pas_grammes=1000, mode="superieur")
DRAP = uuid.uuid4()
TAIE = uuid.uuid4()
COMBO = uuid.uuid4()
ARTICLES = {DRAP: ArticleTarifie("Drap", 1000), TAIE: ArticleTarifie("Taie", 300)}
COMBOS = {COMBO: CombinaisonTarifee("Drap + taie", 500)}


class TestPoidsFacturable:
    @pytest.mark.parametrize(
        "poids, pas, mode, attendu",
        [
            (1, 1000, "superieur", 1000),
            (999, 1000, "superieur", 1000),
            (1000, 1000, "superieur", 1000),
            (1001, 1000, "superieur", 2000),
            (3400, 1000, "superieur", 4000),
            (3400, 500, "superieur", 3500),
            (3401, 100, "superieur", 3500),
            (3401, 1, "superieur", 3401),
            (3449, 100, "proche", 3400),
            (3450, 100, "proche", 3500),
            (3500, 1000, "proche", 4000),  # demi vers le haut
            (3499, 1000, "proche", 3000),
            (1, 1000, "proche", 1000),  # jamais 0 : au moins un pas
            (40, 100, "proche", 100),
            (3400, 1, "proche", 3400),
        ],
    )
    def test_arrondi(self, poids, pas, mode, attendu):
        assert poids_facturable(poids, pas, mode) == attendu

    @pytest.mark.parametrize("poids", [0, -5, None])
    def test_poids_nul_ou_negatif_refuse(self, poids):
        with pytest.raises(PressingRegleError, match="au moins 1 gramme"):
            poids_facturable(poids, 1000, "superieur")

    def test_poids_aberrant_refuse(self):
        with pytest.raises(PressingRegleError):
            poids_facturable(2_000_000, 1000, "superieur")

    def test_pas_et_mode_invalides(self):
        with pytest.raises(PressingRegleError):
            poids_facturable(1000, 250, "superieur")
        with pytest.raises(PressingRegleError):
            poids_facturable(1000, 1000, "arrondi")


class TestCalcul:
    def test_montant_lavage_500_fcfa_le_kilo(self):
        assert montant_au_kilo(1000, 500) == 500
        assert montant_au_kilo(4000, 500) == 2000
        assert montant_au_kilo(3400, 500) == 1700

    def test_montant_arrondi_au_franc_demi_haut(self):
        assert montant_au_kilo(1, 500) == 1  # 0,5 -> 1
        assert montant_au_kilo(3, 500) == 2  # 1,5 -> 2
        assert montant_au_kilo(1, 499) == 0  # 0,499 -> 0

    def test_lavage_seul(self):
        r = calculer(TARIF, DemandeCalcul(poids_grams=3400), {}, {})
        assert r["poids_facturable_grams"] == 4000
        assert r["lignes"] == [
            {
                "type": "lavage",
                "libelle": "Lavage, rinçage et essorage",
                "quantite": Decimal("4.000"),
                "prix_unitaire": 500,
                "montant": 2000,
            }
        ]
        assert (r["sous_total"], r["remise_fcfa"], r["total"], r["devise"]) == (2000, 0, 2000, "FCFA")

    def test_pas_100_proche_facture_le_poids_arrondi(self):
        tarif = Tarif(500, None, 100, "proche")
        r = calculer(tarif, DemandeCalcul(poids_grams=3450), {}, {})
        assert r["poids_facturable_grams"] == 3500
        assert r["lignes"][0]["quantite"] == Decimal("3.500")
        assert r["total"] == 1750

    def test_sechage_sans_tarif_refuse(self):
        with pytest.raises(PressingRegleError, match="séchage"):
            calculer(TARIF, DemandeCalcul(poids_grams=3000, sechage=True), {}, {})

    def test_sechage_facture_au_kilo(self):
        tarif = Tarif(500, 300, 1000, "superieur")
        r = calculer(tarif, DemandeCalcul(poids_grams=3400, sechage=True), {}, {})
        sechage = [l for l in r["lignes"] if l["type"] == "sechage"][0]
        assert (sechage["prix_unitaire"], sechage["montant"]) == (300, 1200)
        assert r["total"] == 2000 + 1200

    def test_sechage_seul_sans_lavage(self):
        tarif = Tarif(500, 300, 1000, "superieur")
        r = calculer(tarif, DemandeCalcul(poids_grams=1000, lavage=False, sechage=True), {}, {})
        assert [l["type"] for l in r["lignes"]] == ["sechage"]

    def test_poids_obligatoire_pour_lavage(self):
        with pytest.raises(PressingRegleError, match="poids est obligatoire"):
            calculer(TARIF, DemandeCalcul(poids_grams=None), {}, {})

    def test_articles_seuls_sans_poids(self):
        r = calculer(
            TARIF, DemandeCalcul(lavage=False, articles=[(DRAP, 3)]), ARTICLES, {}
        )
        assert r["poids_facturable_grams"] is None
        assert r["lignes"][0]["montant"] == 3000 and r["total"] == 3000

    def test_combinaison_facturee_au_prix_fixe_pas_a_la_somme(self):
        # Drap (1000) + taie (300) = 1300 en articles ; la combinaison est fixée à 500.
        r = calculer(
            TARIF,
            DemandeCalcul(lavage=False, combinaisons=[(COMBO, 2)]),
            ARTICLES,
            COMBOS,
        )
        ligne = r["lignes"][0]
        assert (ligne["type"], ligne["prix_unitaire"], ligne["montant"]) == ("combinaison", 500, 1000)
        assert r["total"] == 1000

    def test_exemple_complet(self):
        r = calculer(
            TARIF,
            DemandeCalcul(
                poids_grams=3400, articles=[(DRAP, 1)], combinaisons=[(COMBO, 1)]
            ),
            ARTICLES,
            COMBOS,
        )
        assert [l["montant"] for l in r["lignes"]] == [2000, 1000, 500]
        assert r["sous_total"] == 3500 and r["total"] == 3500

    def test_remise(self):
        r = calculer(TARIF, DemandeCalcul(poids_grams=1000, remise_fcfa=200), {}, {})
        assert (r["sous_total"], r["remise_fcfa"], r["total"]) == (500, 200, 300)
        r = calculer(TARIF, DemandeCalcul(poids_grams=1000, remise_fcfa=500), {}, {})
        assert r["total"] == 0

    def test_remise_superieure_au_sous_total_refusee(self):
        with pytest.raises(PressingRegleError, match="dépasse le sous-total"):
            calculer(TARIF, DemandeCalcul(poids_grams=1000, remise_fcfa=501), {}, {})

    def test_article_desactive_refuse(self):
        inactifs = {DRAP: ArticleTarifie("Drap", 1000, actif=False)}
        with pytest.raises(PressingRegleError, match="désactivé"):
            calculer(TARIF, DemandeCalcul(lavage=False, articles=[(DRAP, 1)]), inactifs, {})

    def test_combinaison_desactivee_refusee(self):
        inactives = {COMBO: CombinaisonTarifee("Drap + taie", 500, actif=False)}
        with pytest.raises(PressingRegleError, match="désactivée"):
            calculer(TARIF, DemandeCalcul(lavage=False, combinaisons=[(COMBO, 1)]), {}, inactives)

    def test_article_inconnu_ou_autre_restaurant_refuse(self):
        with pytest.raises(PressingRegleError, match="introuvable"):
            calculer(TARIF, DemandeCalcul(lavage=False, articles=[(uuid.uuid4(), 1)]), ARTICLES, {})

    def test_aucune_ligne_refusee(self):
        with pytest.raises(PressingRegleError, match="Aucune ligne"):
            calculer(TARIF, DemandeCalcul(lavage=False), {}, {})


# --------------------------------------------------------------------------- #
# Intégration : API
# --------------------------------------------------------------------------- #
@pytest.fixture
def fresh(client):
    """Propriétaire neuf avec son restaurant (l'import est limité à 1/minute et par restaurant)."""

    def make(tag="pressing"):
        user = register(client, tag)
        headers = headers_for(client, user)
        created = client.post("/v1/restaurants", headers=headers, json={"name": f"Pressing {tag}"})
        assert created.status_code == 201, created.text
        return {**user, "headers": headers, "rid": created.json()["id"]}

    return make


def url(owner, path=""):
    return f"/v1/restaurants/{owner['rid']}/pressing{path}"


def make_article(client, owner, nom, prix, **extra):
    response = client.post(
        url(owner, "/articles"), headers=owner["headers"], json={"nom": nom, "prix_unitaire": prix, **extra}
    )
    assert response.status_code == 201, response.text
    return response.json()


def make_combinaison(client, owner, nom, prix, lignes):
    response = client.post(
        url(owner, "/combinaisons"),
        headers=owner["headers"],
        json={"nom": nom, "prix": prix, "lignes": lignes},
    )
    assert response.status_code == 201, response.text
    return response.json()


def audit(engine, owner, action):
    with engine.begin() as conn:
        conn.execute(text("select set_config('app.is_platform_admin', 'true', true)"))
        return conn.execute(
            text(
                "select before_state, after_state, actor_id from ros_audit_logs"
                " where restaurant_id = :rid and action = :action order by created_at, id"
            ),
            {"rid": owner["rid"], "action": action},
        ).all()


class TestTarifs:
    def test_tarif_par_defaut_cree_a_la_premiere_lecture(self, client, fresh):
        owner = fresh()
        response = client.get(url(owner, "/tarifs"), headers=owner["headers"])
        assert response.status_code == 200, response.text
        assert response.json() == {
            "prix_kg": 500,
            "prix_sechage_kg": None,
            "arrondi": {"pas_grammes": 1000, "mode": "superieur"},
            "devise": "FCFA",
        }
        # Deuxième lecture : toujours une seule ligne.
        assert client.get(url(owner, "/tarifs"), headers=owner["headers"]).json()["prix_kg"] == 500

    def test_put_tarif_et_journal(self, client, engine, fresh):
        owner = fresh()
        body = {
            "prix_kg": 600,
            "prix_sechage_kg": 250,
            "arrondi": {"pas_grammes": 100, "mode": "proche"},
            "devise": "FCFA",
        }
        response = client.put(url(owner, "/tarifs"), headers=owner["headers"], json=body)
        assert response.status_code == 200, response.text
        assert response.json() == body
        assert client.get(url(owner, "/tarifs"), headers=owner["headers"]).json() == body

        (avant, apres, acteur), = audit(engine, owner, "PRESSING_TARIF_UPDATE")
        assert avant["prix_kg"] == 500 and avant["prix_sechage_kg"] is None
        assert apres["prix_kg"] == 600 and apres["arrondi"] == {"pas_grammes": 100, "mode": "proche"}
        assert str(acteur) == owner["id"]

        # null = séchage non proposé
        body["prix_sechage_kg"] = None
        assert client.put(url(owner, "/tarifs"), headers=owner["headers"], json=body).json()[
            "prix_sechage_kg"
        ] is None

    @pytest.mark.parametrize(
        "patch",
        [
            {"prix_kg": -1},
            {"prix_kg": 12.5},
            {"prix_sechage_kg": -10},
            {"arrondi": {"pas_grammes": 250, "mode": "superieur"}},
            {"arrondi": {"pas_grammes": 100, "mode": "plafond"}},
        ],
    )
    def test_validation(self, client, fresh, patch):
        owner = fresh()
        body = {
            "prix_kg": 500,
            "prix_sechage_kg": None,
            "arrondi": {"pas_grammes": 1000, "mode": "superieur"},
            **patch,
        }
        response = client.put(url(owner, "/tarifs"), headers=owner["headers"], json=body)
        assert response.status_code == 422, response.text


class TestArticles:
    def test_crud_et_desactivation(self, client, engine, fresh):
        owner = fresh()
        drap = make_article(client, owner, "Drap", 1000, categorie="Linge de lit")
        assert drap["actif"] is True and drap["categorie"] == "Linge de lit"
        assert set(drap) == {"id", "nom", "prix_unitaire", "categorie", "actif"}

        got = client.get(url(owner, f"/articles/{drap['id']}"), headers=owner["headers"])
        assert got.status_code == 200 and got.json() == drap

        put = client.put(
            url(owner, f"/articles/{drap['id']}"),
            headers=owner["headers"],
            json={"nom": "Grand drap", "prix_unitaire": 1200, "categorie": None},
        )
        assert put.status_code == 200, put.text
        assert put.json()["nom"] == "Grand drap" and put.json()["prix_unitaire"] == 1200
        assert put.json()["categorie"] is None and put.json()["actif"] is True
        (_, _), (avant, apres, _) = [
            (r[0], r[1]) for r in audit(engine, owner, "PRESSING_ARTICLE_CREATE")
        ] + audit(engine, owner, "PRESSING_ARTICLE_UPDATE")
        assert avant["prix_unitaire"] == 1000 and apres["prix_unitaire"] == 1200

        deleted = client.delete(url(owner, f"/articles/{drap['id']}"), headers=owner["headers"])
        assert deleted.status_code == 204 and deleted.content == b""
        assert client.get(url(owner, f"/articles/{drap['id']}"), headers=owner["headers"]).json()[
            "actif"
        ] is False
        assert len(audit(engine, owner, "PRESSING_ARTICLE_DEACTIVATE")) == 1

        taie = make_article(client, owner, "Taie", 300)
        listing = lambda **p: [
            a["nom"] for a in client.get(url(owner, "/articles"), headers=owner["headers"], params=p).json()
        ]
        assert listing() == ["Grand drap", "Taie"]
        assert listing(actif="true") == ["Taie"]
        assert listing(actif="false") == ["Grand drap"]
        assert taie["actif"] is True

    def test_doublon_de_nom_409_sans_tenir_compte_de_la_casse(self, client, fresh):
        owner = fresh()
        make_article(client, owner, "Drap", 1000)
        for nom in ("Drap", "  drap ", "DRAP"):
            response = client.post(
                url(owner, "/articles"), headers=owner["headers"], json={"nom": nom, "prix_unitaire": 5}
            )
            assert response.status_code == 409, (nom, response.text)
        assert "existe déjà" in response.json()["detail"]
        autre = make_article(client, owner, "Taie", 300)
        renomme = client.put(
            url(owner, f"/articles/{autre['id']}"),
            headers=owner["headers"],
            json={"nom": "drap", "prix_unitaire": 300},
        )
        assert renomme.status_code == 409

    def test_meme_nom_possible_dans_un_autre_restaurant(self, client, fresh):
        a, b = fresh("a"), fresh("b")
        make_article(client, a, "Drap", 1000)
        make_article(client, b, "Drap", 800)

    @pytest.mark.parametrize(
        "body",
        [
            {"nom": "", "prix_unitaire": 1},
            {"nom": "   ", "prix_unitaire": 1},
            {"nom": "x" * 121, "prix_unitaire": 1},
            {"nom": "Ok", "prix_unitaire": -1},
            {"nom": "Ok"},
        ],
    )
    def test_validation(self, client, fresh, body):
        owner = fresh()
        assert client.post(url(owner, "/articles"), headers=owner["headers"], json=body).status_code == 422

    def test_inconnu_404(self, client, fresh):
        owner = fresh()
        response = client.get(url(owner, f"/articles/{uuid.uuid4()}"), headers=owner["headers"])
        assert response.status_code == 404


class TestCombinaisons:
    def test_crud_et_total_articles(self, client, engine, fresh):
        owner = fresh()
        drap = make_article(client, owner, "Drap", 1000)
        taie = make_article(client, owner, "Taie d'oreiller", 300)
        combo = make_combinaison(
            client, owner, "Drap + taie", 500,
            [{"article_id": drap["id"], "quantite": 1}, {"article_id": taie["id"], "quantite": 2}],
        )  # fmt: skip
        assert combo["prix"] == 500 and combo["actif"] is True
        assert combo["total_articles"] == 1000 + 2 * 300
        assert {l["article_nom"]: (l["quantite"], l["prix_unitaire_reference"]) for l in combo["lignes"]} == {
            "Drap": (1, 1000),
            "Taie d'oreiller": (2, 300),
        }
        assert set(combo) == {"id", "nom", "prix", "actif", "lignes", "total_articles"}
        assert set(combo["lignes"][0]) == {"article_id", "article_nom", "quantite", "prix_unitaire_reference"}

        listing = client.get(url(owner, "/combinaisons"), headers=owner["headers"]).json()
        assert [c["id"] for c in listing] == [combo["id"]]
        assert listing[0]["total_articles"] == 1600

        # PUT remplace les lignes
        put = client.put(
            url(owner, f"/combinaisons/{combo['id']}"),
            headers=owner["headers"],
            json={"nom": "Drap + taie", "prix": 700, "lignes": [{"article_id": drap["id"], "quantite": 1}]},
        )
        assert put.status_code == 200, put.text
        assert put.json()["prix"] == 700 and put.json()["total_articles"] == 1000
        assert [l["article_nom"] for l in put.json()["lignes"]] == ["Drap"]
        got = client.get(url(owner, f"/combinaisons/{combo['id']}"), headers=owner["headers"])
        assert got.json() == put.json()
        (avant, apres, _), = audit(engine, owner, "PRESSING_COMBINAISON_UPDATE")
        assert avant["prix"] == 500 and apres["prix"] == 700 and len(apres["lignes"]) == 1

        # Doublon de nom
        dup = client.post(
            url(owner, "/combinaisons"),
            headers=owner["headers"],
            json={"nom": "DRAP + TAIE", "prix": 1, "lignes": [{"article_id": drap["id"], "quantite": 1}]},
        )
        assert dup.status_code == 409

        # Désactivation
        deleted = client.delete(url(owner, f"/combinaisons/{combo['id']}"), headers=owner["headers"])
        assert deleted.status_code == 204
        assert client.get(url(owner, f"/combinaisons/{combo['id']}"), headers=owner["headers"]).json()[
            "actif"
        ] is False

    def test_regles_sur_les_lignes(self, client, fresh):
        owner = fresh()
        drap = make_article(client, owner, "Drap", 1000)
        vieux = make_article(client, owner, "Vieux", 100, actif=False)
        post = lambda lignes: client.post(
            url(owner, "/combinaisons"),
            headers=owner["headers"],
            json={"nom": f"C{uuid.uuid4().hex[:6]}", "prix": 10, "lignes": lignes},
        )
        assert post([]).status_code == 422  # au moins une ligne
        assert post([{"article_id": drap["id"], "quantite": 0}]).status_code == 422
        assert post([{"article_id": str(uuid.uuid4()), "quantite": 1}]).status_code == 422
        inactif = post([{"article_id": vieux["id"], "quantite": 1}])
        assert inactif.status_code == 422 and "désactivé" in inactif.json()["detail"]
        double = post([{"article_id": drap["id"], "quantite": 1}, {"article_id": drap["id"], "quantite": 2}])
        assert double.status_code == 422

    def test_article_desactive_reste_lisible_dans_une_combinaison_existante(self, client, fresh):
        owner = fresh()
        drap = make_article(client, owner, "Drap", 1000)
        combo = make_combinaison(client, owner, "Solo", 400, [{"article_id": drap["id"], "quantite": 1}])
        assert client.delete(url(owner, f"/articles/{drap['id']}"), headers=owner["headers"]).status_code == 204
        got = client.get(url(owner, f"/combinaisons/{combo['id']}"), headers=owner["headers"]).json()
        assert got["lignes"][0]["article_nom"] == "Drap" and got["total_articles"] == 1000
        # Un nouveau calcul refuse l'article, mais la combinaison reste facturable.
        refus = client.post(
            url(owner, "/calcul"), headers=owner["headers"],
            json={"lavage": False, "articles": [{"article_id": drap["id"], "quantite": 1}]},
        )  # fmt: skip
        assert refus.status_code == 422
        ok = client.post(
            url(owner, "/calcul"), headers=owner["headers"],
            json={"lavage": False, "combinaisons": [{"combinaison_id": combo["id"], "quantite": 1}]},
        )  # fmt: skip
        assert ok.status_code == 200 and ok.json()["total"] == 400


class TestCalcul:
    def test_bout_en_bout(self, client, fresh):
        owner = fresh()
        drap = make_article(client, owner, "Drap", 1000)
        taie = make_article(client, owner, "Taie", 300)
        combo = make_combinaison(
            client, owner, "Drap + taie", 500,
            [{"article_id": drap["id"], "quantite": 1}, {"article_id": taie["id"], "quantite": 1}],
        )  # fmt: skip
        response = client.post(
            url(owner, "/calcul"),
            headers=owner["headers"],
            json={
                "poids_grams": 3400,
                "articles": [{"article_id": drap["id"], "quantite": 1}],
                "combinaisons": [{"combinaison_id": combo["id"], "quantite": 1}],
            },
        )
        assert response.status_code == 200, response.text
        assert response.json() == {
            "poids_facturable_grams": 4000,
            "lignes": [
                {"type": "lavage", "libelle": "Lavage, rinçage et essorage", "quantite": "4.000",
                 "prix_unitaire": 500, "montant": 2000},
                {"type": "article", "libelle": "Drap", "quantite": "1", "prix_unitaire": 1000, "montant": 1000},
                {"type": "combinaison", "libelle": "Drap + taie", "quantite": "1", "prix_unitaire": 500,
                 "montant": 500},
            ],  # fmt: skip
            "sous_total": 3500,
            "remise_fcfa": 0,
            "total": 3500,
            "devise": "FCFA",
        }

    def test_tarif_modifie_et_remise(self, client, fresh):
        owner = fresh()
        client.put(
            url(owner, "/tarifs"),
            headers=owner["headers"],
            json={"prix_kg": 600, "prix_sechage_kg": 200, "arrondi": {"pas_grammes": 100, "mode": "proche"}},
        )
        response = client.post(
            url(owner, "/calcul"),
            headers=owner["headers"],
            json={"poids_grams": 3450, "sechage": True, "remise_fcfa": 100},
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["poids_facturable_grams"] == 3500
        assert [l["montant"] for l in body["lignes"]] == [2100, 700]
        assert (body["sous_total"], body["remise_fcfa"], body["total"]) == (2800, 100, 2700)

    def test_calcul_n_ecrit_rien(self, client, engine, fresh):
        owner = fresh()
        client.post(url(owner, "/calcul"), headers=owner["headers"], json={"poids_grams": 1000})
        assert audit(engine, owner, "PRESSING_TARIF_UPDATE") == []

    @pytest.mark.parametrize(
        "corps, extrait",
        [
            ({"poids_grams": 1000, "sechage": True}, "séchage"),
            ({"poids_grams": None}, "poids est obligatoire"),
            ({"poids_grams": 1000, "remise_fcfa": 501}, "dépasse le sous-total"),
            ({"lavage": False}, "Aucune ligne"),
        ],
    )
    def test_regles_metier_422(self, client, fresh, corps, extrait):
        owner = fresh()
        response = client.post(url(owner, "/calcul"), headers=owner["headers"], json=corps)
        assert response.status_code == 422, response.text
        assert extrait in response.json()["detail"]

    @pytest.mark.parametrize("corps", [{"poids_grams": 0}, {"poids_grams": 1000, "remise_fcfa": -1}])
    def test_validation_du_corps_422(self, client, fresh, corps):
        owner = fresh()
        assert client.post(url(owner, "/calcul"), headers=owner["headers"], json=corps).status_code == 422

    def test_isolation_entre_deux_restaurants(self, client, fresh):
        a, b = fresh("iso-a"), fresh("iso-b")
        drap = make_article(client, a, "Drap", 1000)
        combo = make_combinaison(client, a, "Solo", 400, [{"article_id": drap["id"], "quantite": 1}])
        # Lecture directe
        assert client.get(url(b, f"/articles/{drap['id']}"), headers=b["headers"]).status_code == 404
        assert client.get(url(b, f"/combinaisons/{combo['id']}"), headers=b["headers"]).status_code == 404
        assert client.get(url(b, "/articles"), headers=b["headers"]).json() == []
        # Passer par le restaurant A avec le jeton de B : refusé
        assert client.get(url(a, "/articles"), headers=b["headers"]).status_code in (403, 404)
        # Calcul et composition
        calc_article = client.post(
            url(b, "/calcul"), headers=b["headers"],
            json={"lavage": False, "articles": [{"article_id": drap["id"], "quantite": 1}]},
        )  # fmt: skip
        assert calc_article.status_code == 422 and "introuvable" in calc_article.json()["detail"]
        calc_combo = client.post(
            url(b, "/calcul"), headers=b["headers"],
            json={"lavage": False, "combinaisons": [{"combinaison_id": combo["id"], "quantite": 1}]},
        )  # fmt: skip
        assert calc_combo.status_code == 422
        compose = client.post(
            url(b, "/combinaisons"), headers=b["headers"],
            json={"nom": "Vol", "prix": 1, "lignes": [{"article_id": drap["id"], "quantite": 1}]},
        )  # fmt: skip
        assert compose.status_code == 422
        # Le tarif de B n'est pas celui de A
        client.put(
            url(a, "/tarifs"), headers=a["headers"],
            json={"prix_kg": 900, "prix_sechage_kg": None, "arrondi": {"pas_grammes": 1, "mode": "proche"}},
        )  # fmt: skip
        assert client.get(url(b, "/tarifs"), headers=b["headers"]).json()["prix_kg"] == 500


class TestDroits:
    def test_roles(self, client, engine, fresh):
        owner = fresh()
        make_article(client, owner, "Drap", 1000)
        serveur = make_member(client, engine, owner, "waiter")
        caissier = make_member(client, engine, owner, "cashier")
        chef = make_member(client, engine, owner, "chef")
        manager = make_member(client, engine, owner, "manager")

        for lecteur in (serveur, caissier, manager):
            assert client.get(url(owner, "/tarifs"), headers=lecteur["headers"]).status_code == 200
            assert client.get(url(owner, "/articles"), headers=lecteur["headers"]).status_code == 200
            assert client.get(url(owner, "/combinaisons"), headers=lecteur["headers"]).status_code == 200
            assert (
                client.post(url(owner, "/calcul"), headers=lecteur["headers"], json={"poids_grams": 1000}).status_code
                == 200
            )
        assert client.get(url(owner, "/tarifs"), headers=chef["headers"]).status_code == 403

        tarif = {"prix_kg": 700, "prix_sechage_kg": None, "arrondi": {"pas_grammes": 1000, "mode": "superieur"}}
        for ecrivain in (serveur, caissier, chef):
            h = ecrivain["headers"]
            assert client.put(url(owner, "/tarifs"), headers=h, json=tarif).status_code == 403
            assert (
                client.post(url(owner, "/articles"), headers=h, json={"nom": "X", "prix_unitaire": 1}).status_code
                == 403
            )
            assert client.get(url(owner, "/modele"), headers=h, params={"type": "articles"}).status_code == 403
        assert client.put(url(owner, "/tarifs"), headers=manager["headers"], json=tarif).status_code == 200

    def test_sans_jeton_401(self, client, fresh):
        owner = fresh()
        assert client.get(url(owner, "/tarifs")).status_code in (401, 403)


# --------------------------------------------------------------------------- #
# Modèle et import
# --------------------------------------------------------------------------- #
def workbook_bytes(header, rows, title="Articles"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(header)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def csv_bytes(*lines):
    return "\r\n".join(lines).encode("utf-8-sig")


def clear_import_limit(engine, owner):
    with engine.begin() as conn:
        conn.execute(text("select set_config('app.is_platform_admin', 'true', true)"))
        conn.execute(
            text("delete from ros_audit_logs where restaurant_id = :rid and action = 'PRESSING_IMPORT'"),
            {"rid": owner["rid"]},
        )


def post_import(client, owner, kind, content, name="fichier.xlsx", mime=XLSX):
    return client.post(
        url(owner, "/import"), headers=owner["headers"], params={"type": kind},
        files={"file": (name, content, mime)},
    )  # fmt: skip


class TestModeleEtImport:
    @pytest.mark.parametrize(
        "kind, entetes",
        [
            ("articles", ["nom", "prix_unitaire", "categorie"]),
            (
                "combinaisons",
                ["nom", "prix"] + [c for i in range(1, 6) for c in (f"article_{i}", f"quantite_{i}")],
            ),
        ],
    )
    def test_modele_vierge(self, client, fresh, kind, entetes):
        owner = fresh()
        response = client.get(url(owner, "/modele"), headers=owner["headers"], params={"type": kind})
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(XLSX)
        rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
        assert rows == [tuple(entetes)]

    def test_modele_type_inconnu(self, client, fresh):
        owner = fresh()
        assert client.get(url(owner, "/modele"), headers=owner["headers"], params={"type": "x"}).status_code == 422

    def test_import_articles_xlsx_idempotent(self, client, engine, fresh):
        owner = fresh()
        content = workbook_bytes(
            ["nom", "prix_unitaire", "categorie"],
            [["Drap", 1000, "Linge de lit"], ["Taie", 300, None], ["Couette", 2500, "Linge de lit"]],
        )
        response = post_import(client, owner, "articles", content)
        assert response.status_code == 200, response.text
        assert response.json() == {"imported": 3, "created": 3, "updated": 0, "errors": []}

        # Un second import dans la minute est refusé ...
        again = post_import(client, owner, "articles", content)
        assert again.status_code == 429 and "Retry-After" in again.headers
        # ... et, la limite levée, il met à jour sans dupliquer (casse ignorée).
        clear_import_limit(engine, owner)
        maj = workbook_bytes(["nom", "prix_unitaire", "categorie"], [["DRAP", 1100, None], ["Nappe", 700, None]])
        response = post_import(client, owner, "articles", maj)
        assert response.json() == {"imported": 2, "created": 1, "updated": 1, "errors": []}
        articles = {a["nom"]: a for a in client.get(url(owner, "/articles"), headers=owner["headers"]).json()}
        assert set(articles) == {"Drap", "Taie", "Couette", "Nappe"}
        assert articles["Drap"]["prix_unitaire"] == 1100 and articles["Drap"]["categorie"] == "Linge de lit"

    def test_import_csv_partiel_avec_erreurs(self, client, fresh):
        owner = fresh()
        content = csv_bytes(
            "nom;prix_unitaire;categorie",
            ";500;",  # 2 : nom manquant
            "Rideau;abc;",  # 3 : prix illisible
            "Tapis;-5;",  # 4 : prix négatif
            "Veste;12,5;",  # 5 : pas un entier
            "Pantalon;800;Vêtements",  # 6 : valide
        )
        response = post_import(client, owner, "articles", content, name="a.csv", mime=CSV)
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["imported"] == 1 and body["created"] == 1
        assert {(e["row"], e["field"]) for e in body["errors"]} == {
            (2, "nom"), (3, "prix_unitaire"), (4, "prix_unitaire"), (5, "prix_unitaire"),
        }  # fmt: skip
        assert all(e["message"] for e in body["errors"])
        assert [a["nom"] for a in client.get(url(owner, "/articles"), headers=owner["headers"]).json()] == ["Pantalon"]

    def test_import_combinaisons(self, client, engine, fresh):
        owner = fresh()
        make_article(client, owner, "Drap", 1000)
        make_article(client, owner, "Taie", 300)
        make_article(client, owner, "Vieux", 100, actif=False)
        header = ["nom", "prix", "article_1", "quantite_1", "article_2", "quantite_2"]
        content = workbook_bytes(
            header,
            [
                ["Drap + taie", 500, "drap", 1, "Taie", 2],
                ["Sans article", 100, None, None, None, None],
                ["Article inconnu", 100, "Fantôme", 1, None, None],
                ["Article inactif", 100, "Vieux", 1, None, None],
                ["Quantité par défaut", 900, "Drap", None, None, None],
            ],
            title="Combinaisons",
        )
        response = post_import(client, owner, "combinaisons", content)
        assert response.status_code == 200, response.text
        body = response.json()
        assert (body["imported"], body["created"], body["updated"]) == (2, 2, 0)
        assert {(e["row"], e["field"]) for e in body["errors"]} == {
            (3, "article_1"), (4, "article_1"), (5, "article_1"),
        }  # fmt: skip
        combos = {c["nom"]: c for c in client.get(url(owner, "/combinaisons"), headers=owner["headers"]).json()}
        assert combos["Drap + taie"]["prix"] == 500 and combos["Drap + taie"]["total_articles"] == 1600
        assert combos["Quantité par défaut"]["lignes"][0]["quantite"] == 1

        # Idempotent : même nom -> mise à jour du prix et des lignes
        clear_import_limit(engine, owner)
        maj = workbook_bytes(header, [["drap + TAIE", 650, "Drap", 1, None, None]], title="Combinaisons")
        response = post_import(client, owner, "combinaisons", maj)
        assert response.json() == {"imported": 1, "created": 0, "updated": 1, "errors": []}
        combos = {c["nom"]: c for c in client.get(url(owner, "/combinaisons"), headers=owner["headers"]).json()}
        assert len(combos) == 2
        assert combos["Drap + taie"]["prix"] == 650 and len(combos["Drap + taie"]["lignes"]) == 1

    def test_fichier_inutilisable_400(self, client, fresh):
        owner = fresh()
        assert post_import(client, owner, "articles", b"pas un classeur").status_code == 400
        mauvais_sep = csv_bytes("nom,prix_unitaire", "Drap,100")
        assert post_import(client, owner, "articles", mauvais_sep, name="a.csv", mime=CSV).status_code == 400
        sans_colonne = csv_bytes("nom;categorie", "Drap;x")
        assert post_import(client, owner, "articles", sans_colonne, name="a.csv", mime=CSV).status_code == 400
        assert post_import(client, owner, "articles", b"x", name="a.txt", mime="text/plain").status_code == 400

    def test_trop_gros_413(self, client, fresh):
        owner = fresh()
        gros = b"nom;prix_unitaire\n" + b"a;1\n" * (1_400_000)
        assert post_import(client, owner, "articles", gros, name="a.csv", mime=CSV).status_code == 413

    def test_import_reserve_au_management(self, client, engine, fresh):
        owner = fresh()
        serveur = make_member(client, engine, owner, "waiter")
        content = workbook_bytes(["nom", "prix_unitaire"], [["Drap", 1]])
        response = client.post(
            url(owner, "/import"), headers=serveur["headers"], params={"type": "articles"},
            files={"file": ("a.xlsx", content, XLSX)},
        )  # fmt: skip
        assert response.status_code == 403
